"""Licence-filtered exports of a run's plan (Section 7.4, 11): GeoJSON, XLSX, PDF.

What goes out: our own derived outputs (sites, bundles, costs, served energy, scores,
economics) and an attribution sheet for the sources behind them. What never goes out:
content from RESTRICTED_GOOGLE or RESTRICTED_COMMERCIAL sources (Google Solar results,
DISCOM asset data, operator feed records). Those sources are named in the attribution
sheet only as "used in the analysis; content not exported". Every export carries the
candidate-site label and the data caveats (synthetic registrations, placeholders).
"""

from __future__ import annotations

import io
import json
import uuid
from datetime import UTC, datetime
from typing import Any

from sqlalchemy import func, select, text
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.models.planning import (
    CandidateSite,
    Optimisation,
    OptimisationSite,
    PlanningRun,
    RunSite,
    Scenario,
    SiteConfig,
)

EXPORTABLE = {"OPEN", "GOVERNMENT", "SYNTHETIC"}
LABEL = "Candidate sites: proposals from open data, each requires field verification."


class ExportError(RuntimeError):
    pass


async def plan_data(
    session: AsyncSession, run_id: uuid.UUID, optimisation_id: uuid.UUID | None = None
) -> dict[str, Any]:
    run = await session.get(PlanningRun, run_id)
    if run is None:
        raise ExportError("Run not found")
    scenario = await session.get(Scenario, run.scenario_id)
    assert scenario is not None
    opt_id = optimisation_id or uuid.UUID(
        (run.funnel.get("optimisation") or {}).get("base_id") or str(uuid.UUID(int=0))
    )
    opt = await session.get(Optimisation, opt_id)
    if opt is None or opt.run_id != run_id:
        raise ExportError("This run has no optimised plan to export")
    rows = (
        await session.execute(
            select(
                OptimisationSite,
                CandidateSite.name,
                CandidateSite.host_type,
                CandidateSite.origin,
                func.ST_Y(CandidateSite.geom),
                func.ST_X(CandidateSite.geom),
                RunSite.rank,
                RunSite.score_total,
            )
            .join(CandidateSite, CandidateSite.id == OptimisationSite.site_id)
            .outerjoin(
                RunSite, (RunSite.site_id == OptimisationSite.site_id) & (RunSite.run_id == run_id)
            )
            .where(OptimisationSite.optimisation_id == opt.id)
            .order_by(OptimisationSite.served_kwh.desc())
        )
    ).all()
    configs = {
        c.site_id: c
        for c in (
            await session.execute(
                select(SiteConfig).where(
                    SiteConfig.optimisation_id == opt.id, SiteConfig.subsidy.is_(False)
                )
            )
        ).scalars()
    }
    sites = []
    for s, name, host, origin, lat, lng, rank, score in rows:
        cfg = configs.get(s.site_id)
        fin = cfg.finance if cfg else None
        sites.append(
            {
                "site_id": str(s.site_id),
                "name": name,
                "host_type": host,
                "origin": origin,
                "lat": round(float(lat), 6),
                "lng": round(float(lng), 6),
                "rank": rank,
                "score_total": score,
                "bundle": s.bundle,
                "charge_points": s.charge_points,
                "cost_inr": s.cost_inr,
                "capacity_kwh_per_day": s.capacity_kwh,
                "served_kwh_per_day": s.served_kwh,
                "utilisation": round(s.served_kwh / s.capacity_kwh, 3) if s.capacity_kwh else None,
                "phase_year": s.detail.get("phase_year"),
                "chargers_sized": cfg.sizing.get("chargers") if cfg else None,
                "connection": cfg.sizing.get("electrical", {}).get("connection") if cfg else None,
                "capex_sized_inr": fin.get("capex_total") if fin else None,
                "npv_p10_inr": _mc(fin, "npv", "p10"),
                "npv_p50_inr": _mc(fin, "npv", "p50"),
                "npv_p90_inr": _mc(fin, "npv", "p90"),
                "irr_p50": _mc(fin, "irr", "p50"),
                # energy results may include Google Solar content: never exported
            }
        )
    return {
        "run": run,
        "scenario": scenario,
        "optimisation": opt,
        "sites": sites,
        "attribution": await attribution(session, run_id),
        "caveats": caveats(run, opt),
        "generated_at": datetime.now(UTC).isoformat(timespec="seconds"),
    }


def _mc(fin: dict[str, Any] | None, metric: str, q: str) -> float | None:
    if not fin:
        return None
    v = (fin.get("monte_carlo") or {}).get(metric, {}).get(q)
    return None if v is None else round(float(v), 4 if metric == "irr" else 0)


async def attribution(session: AsyncSession, run_id: uuid.UUID) -> list[dict[str, Any]]:
    """Sources behind the run's evidence, with licence and retrieval date."""
    rows = (
        await session.execute(
            text(
                "select ds.id, ds.name, ds.license_class, ds.attribution, ds.license_text, "
                "max(sn.retrieved_at) from evidence e "
                "cross join lateral unnest(e.snapshot_ids) sid "
                "join dataset_snapshot sn on sn.id = sid "
                "join data_source ds on ds.id = sn.source_id "
                "where e.run_id = :run group by 1, 2, 3, 4, 5 order by 1"
            ),
            {"run": run_id},
        )
    ).all()
    out = []
    for sid, name, lic, attr, lic_text, retrieved in rows:
        exported = lic in EXPORTABLE
        out.append(
            {
                "source": name if exported else "Contractual/restricted source",
                "source_id": sid if exported else None,
                "license_class": lic,
                "attribution": attr if exported else None,
                "license": lic_text if exported else None,
                "retrieved": retrieved.date().isoformat() if retrieved else None,
                "note": "" if exported else "Used in the analysis; content not exported",
            }
        )
    return out


def caveats(run: PlanningRun, opt: Optimisation) -> list[str]:
    out = [LABEL]
    if run.inputs.get("demand_synthetic"):
        out.append(
            "EV registrations are a SYNTHETIC stand-in (no VAHAN export loaded): demand, gap "
            "and served-energy figures illustrate the method and are not estimates."
        )
    placeholders = run.inputs.get("placeholders_in_use") or opt.params.get("placeholders") or []
    if placeholders:
        out.append(
            f"{len(placeholders)} assumptions are illustrative placeholders (costs, tariffs, "
            "thresholds): money figures have confidence NONE until they are replaced."
        )
    out.append("Grid capacity is a proximity estimate unless DISCOM data is loaded.")
    return out


# --- formats ---------------------------------------------------------------------------


def to_geojson(d: dict[str, Any]) -> bytes:
    fc = {
        "type": "FeatureCollection",
        "name": d["scenario"].name,
        "label": LABEL,
        "generated_at": d["generated_at"],
        "caveats": d["caveats"],
        "attribution": [a for a in d["attribution"] if a["license_class"] in EXPORTABLE],
        "features": [
            {
                "type": "Feature",
                "geometry": {"type": "Point", "coordinates": [s["lng"], s["lat"]]},
                "properties": {k: v for k, v in s.items() if k not in ("lat", "lng")},
            }
            for s in d["sites"]
        ],
    }
    return json.dumps(fc, ensure_ascii=False, indent=1, default=str).encode()


def to_xlsx(d: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Plan"
    cols = list(d["sites"][0].keys()) if d["sites"] else ["site_id"]
    ws.append([LABEL])
    ws.append(cols)
    for c in ws[2]:
        c.font = Font(bold=True)
    for s in d["sites"]:
        ws.append([s.get(c) for c in cols])
    sc, opt = d["scenario"], d["optimisation"]
    summary = wb.create_sheet("Scenario")
    for k, v in [
        ("Scenario", sc.name),
        ("Region", sc.region_id),
        ("Target year", sc.target_year),
        ("Adoption", f"{sc.adoption_case} x{sc.adoption_multiplier:g}"),
        ("Charger classes", ", ".join(sc.charger_classes)),
        ("Budget (INR)", sc.budget_inr),
        ("Plan", opt.params.get("label") or opt.kind),
        ("Solver status", opt.status),
        *[(f"KPI: {k}", v) for k, v in opt.kpis.items() if isinstance(v, int | float)],
        ("Generated", d["generated_at"]),
    ]:
        summary.append([k, v])
    funnel = wb.create_sheet("Funnel")
    for k, v in (d["run"].funnel or {}).items():
        funnel.append([k, v.get("total") if isinstance(v, dict) else v])
    cav = wb.create_sheet("Caveats")
    for c in d["caveats"]:
        cav.append([c])
    attr = wb.create_sheet("Attribution")
    attr.append(["Source", "Licence class", "Attribution", "Licence", "Retrieved", "Note"])
    for c in attr[1]:
        c.font = Font(bold=True)
    for a in d["attribution"]:
        attr.append(
            [
                a["source"],
                a["license_class"],
                a["attribution"],
                a["license"],
                a["retrieved"],
                a["note"],
            ]
        )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(d: dict[str, Any]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=f"ChargeGrid plan: {d['scenario'].name}",
    )
    st = getSampleStyleSheet()
    sc, opt = d["scenario"], d["optimisation"]
    # The built-in PDF fonts have no rupee glyph.
    title = sc.name.replace("₹", "INR ")
    k = opt.kpis
    story: list[Any] = [
        Paragraph(f"ChargeGrid AI plan: {title}", st["Title"]),
        Paragraph(
            f"{sc.region_id} · target {sc.target_year} · {', '.join(sc.charger_classes)}"
            f" · generated {d['generated_at']}",
            st["Normal"],
        ),
        Spacer(1, 4 * mm),
        Paragraph(f"<b>{LABEL}</b>", st["Normal"]),
        Spacer(1, 3 * mm),
        Paragraph(
            f"Sites: {k.get('n_sites')} · charge points: {k.get('n_charge_points')} · "
            f"plan cost: INR {k.get('cost_inr'):,.0f} · served: "
            f"{k.get('coverage_kwh'):,.0f} kWh/day of {k.get('demand_kwh'):,.0f} kWh/day "
            f"unmet demand"
            if k.get("cost_inr") is not None
            and k.get("coverage_kwh") is not None
            and k.get("demand_kwh") is not None
            else "KPIs unavailable",
            st["Normal"],
        ),
        Spacer(1, 4 * mm),
    ]
    header = [
        "#",
        "Site",
        "Host",
        "Bundle",
        "CPs",
        "Cost (INR L)",
        "Served kWh/d",
        "Util.",
        "NPV P50 (INR L)",
    ]
    data = [header] + [
        [
            i + 1,
            (s["name"] or "Unnamed")[:40],
            (s["host_type"] or "").replace("_", " ").lower(),
            s["bundle"],
            s["charge_points"],
            f"{s['cost_inr'] / 1e5:,.1f}",
            f"{s['served_kwh_per_day']:,.0f}",
            f"{s['utilisation']:.0%}" if s["utilisation"] is not None else "",
            f"{s['npv_p50_inr'] / 1e5:,.1f}" if s["npv_p50_inr"] is not None else "",
        ]
        for i, s in enumerate(d["sites"])
    ]
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8),
                ("FONT", (0, 1), (-1, -1), "Helvetica", 8),
                ("LINEBELOW", (0, 0), (-1, 0), 0.5, colors.black),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f1ec")]),
                ("ALIGN", (4, 1), (-1, -1), "RIGHT"),
            ]
        )
    )
    story += [table, Spacer(1, 5 * mm), Paragraph("Caveats", st["Heading3"])]
    story += [Paragraph(f"• {c}", st["Normal"]) for c in d["caveats"]]
    story += [Spacer(1, 4 * mm), Paragraph("Sources and attribution", st["Heading3"])]
    for a in d["attribution"]:
        parts = [a["license_class"]]
        if a["attribution"]:
            parts.append(a["attribution"])
        if a["retrieved"]:
            parts.append(f"retrieved {a['retrieved']}")
        story.append(Paragraph(f"• {a['source']} ({', '.join(parts)}) {a['note']}", st["Normal"]))
    doc.build(story)
    return buf.getvalue()
